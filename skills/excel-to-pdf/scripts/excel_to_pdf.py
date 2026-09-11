#!/usr/bin/env python3
"""
excel_to_pdf.py —— 用本机 Microsoft Excel 把 .xlsx/.xlsm 转成排版正确的 PDF。

解决两个 Excel 默认导出会踩的坑：
  1. 默认打印设置常常把宽表格拆成几十页 —— 通过直接改写 xlsx 内部 XML
     （sheetPr/pageSetup）强制"横向 + 缩放到一页宽"，而不是依赖
     AppleScript 设置 page setup object（很多属性在自动化接口里设不进去）。
  2. Excel 自动分页经常正好切在纵向合并单元格中间，导致同一行组的内容
     被拆到两页，视觉上看起来像是"断开了"。通过读取 openpyxl 里的合并单元格
     范围，再用 AppleScript 读出 Excel 实际计算出的分页行号，把落在合并区间
     内部的分页点强制挪到该区间的第一行，迭代到没有分页点切在行组中间为止。

只在 macOS + 已安装 Microsoft Excel 的环境下工作。不使用 LibreOffice /
soffice --headless（很多机器没装），也不直接调用 Excel「另存为 PDF」的默认
设置（会产出页数离谱的结果）。

用法：
    python3 excel_to_pdf.py 输入.xlsx [-o 输出.pdf]
        [--paper a4|a3|a2|letter|legal] [--orientation landscape|portrait]
        [--extra-break "工作表名:行号" ...] [--no-render] [--render-dir DIR]

安全性：脚本只操作自己新建的临时工作簿副本，从不改动传入的原始文件，
也不会关闭或影响用户在 Excel 里已经打开的其他工作簿。
"""
from __future__ import annotations

import argparse
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

PAPER_SIZES = {
    "letter": 1,
    "legal": 5,
    "a3": 8,
    "a4": 9,
    "a5": 11,
    "a2": 66,
}

# ---------------------------------------------------------------------------
# AppleScript 模板。所有模板都用工作簿的文件名（不是 "active workbook"）来
# 定位目标，这样即使用户 Excel 里还开着别的文件，也不会误操作到那些文件。
# ---------------------------------------------------------------------------

AS_OPEN = """
on run argv
	set targetPath to item 1 of argv
	with timeout of 600 seconds
		tell application "Microsoft Excel"
			activate
			open (POSIX file targetPath)
		end tell
	end timeout
end run
"""

AS_IS_OPEN = """
on run argv
	set wbName to item 1 of argv
	tell application "Microsoft Excel"
		set allNames to name of every workbook
		if allNames contains wbName then
			return "YES"
		else
			return "NO"
		end if
	end tell
end run
"""

AS_ENUM_BREAKS = """
on run argv
	set wbName to item 1 of argv
	set shName to item 2 of argv
	with timeout of 600 seconds
		tell application "Microsoft Excel"
			set sh to worksheet shName of workbook wbName
			activate object sh
			set out to ""
			repeat with i from 1 to 200
				try
					set out to out & (get first row index of (location of horizontal page break i of sh)) & ","
				on error
					exit repeat
				end try
			end repeat
			return out
		end tell
	end timeout
end run
"""

AS_SET_BREAK = """
on run argv
	set wbName to item 1 of argv
	set shName to item 2 of argv
	set rowN to (item 3 of argv) as integer
	with timeout of 600 seconds
		tell application "Microsoft Excel"
			set sh to worksheet shName of workbook wbName
			activate object sh
			set page break of row rowN of sh to page break manual
			return "OK"
		end tell
	end timeout
end run
"""

AS_EXPORT = """
on run argv
	set wbName to item 1 of argv
	set outPath to item 2 of argv
	with timeout of 600 seconds
		tell application "Microsoft Excel"
			set wb to workbook wbName
			select (every worksheet of wb)
			save workbook as wb filename outPath file format PDF file format
			return "OK"
		end tell
	end timeout
end run
"""

AS_CLOSE = """
on run argv
	set wbName to item 1 of argv
	with timeout of 120 seconds
		tell application "Microsoft Excel"
			try
				close workbook wbName saving no
			end try
		end tell
	end timeout
end run
"""

RENDER_SWIFT = """
import Foundation
import Quartz
let url = URL(fileURLWithPath: CommandLine.arguments[1])
let outDir = CommandLine.arguments[2]
guard let doc = PDFDocument(url: url) else { print("ERROR: cannot open pdf"); exit(1) }
for i in 0..<doc.pageCount {
    guard let p = doc.page(at: i) else { continue }
    let r = p.bounds(for: .mediaBox)
    let scale: CGFloat = 1.5
    let img = NSImage(size: NSSize(width: r.width*scale, height: r.height*scale))
    img.lockFocus()
    NSColor.white.setFill()
    NSRect(x:0,y:0,width:r.width*scale,height:r.height*scale).fill()
    let ctx = NSGraphicsContext.current!.cgContext
    ctx.scaleBy(x: scale, y: scale)
    p.draw(with: .mediaBox, to: ctx)
    img.unlockFocus()
    let rep = NSBitmapImageRep(data: img.tiffRepresentation!)!
    let png = rep.representation(using: .png, properties: [:])!
    try! png.write(to: URL(fileURLWithPath: "\\(outDir)/page\\(i+1).png"))
}
print(doc.pageCount)
"""


def run_osa(work_dir: Path, name: str, script: str, *args: str, timeout: int = 630) -> str:
    script_path = work_dir / f"{name}.applescript"
    script_path.write_text(script, encoding="utf-8")
    proc = subprocess.run(
        ["osascript", str(script_path), *args],
        capture_output=True, text=True, timeout=timeout,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"AppleScript[{name}] 失败: {proc.stderr.strip() or proc.stdout.strip()}")
    return proc.stdout.strip()


def patch_print_settings(src: Path, dst: Path, paper_size: int, orientation: str) -> None:
    """在 xlsx 内部 XML 里为每个 worksheet 注入『缩放到一页宽 + 横/纵向』的打印设置。"""
    zin = zipfile.ZipFile(src)
    sheet_names = [n for n in zin.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n)]
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    page_setup = (
        f'<pageSetup paperSize="{paper_size}" scale="100" orientation="{orientation}" '
        f'fitToWidth="1" fitToHeight="0" horizontalDpi="300" verticalDpi="300"/>'
    )
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheet_names:
            s = data.decode("utf-8")
            # sheetPr / pageSetUpPr fitToPage="1" —— 告诉 Excel 用 pageSetup 的缩放而不是手动分页
            if "<sheetPr" in s:
                m = re.search(r"<sheetPr\b[^>]*?(/?)>", s)
                tag = m.group(0)
                if tag.endswith("/>"):
                    s = s[: m.start()] + tag[:-2] + "><pageSetUpPr fitToPage=\"1\"/></sheetPr>" + s[m.end():]
                elif "<pageSetUpPr" in s:
                    s = re.sub(r"<pageSetUpPr\b[^>]*/>", '<pageSetUpPr fitToPage="1"/>', s, count=1)
                else:
                    s = s[: m.end()] + '<pageSetUpPr fitToPage="1"/>' + s[m.end():]
            else:
                s = re.sub(r"(<worksheet\b[^>]*>)", r'\1<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>', s, count=1)
            # pageSetup 元素本身
            if "<pageSetup" in s:
                s = re.sub(r"<pageSetup\b[^>]*/>", page_setup, s, count=1)
            elif "<pageMargins" in s:
                m2 = re.search(r"<pageMargins\b[^>]*/>", s)
                s = s[: m2.end()] + page_setup + s[m2.end():]
            else:
                s = s.replace(
                    "</worksheet>",
                    '<pageMargins left="0.3" right="0.3" top="0.4" bottom="0.4" header="0.2" footer="0.2"/>'
                    + page_setup + "</worksheet>",
                )
            data = s.encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    zin.close()


def get_vertical_merges(xlsx_path: Path) -> dict[str, list[tuple[int, int]]]:
    """每个 sheet 里跨多行的合并单元格范围 (min_row, max_row)，用于判断分页点是否切错地方。"""
    import openpyxl  # 延迟 import，缺失时给出更友好的报错

    wb = openpyxl.load_workbook(xlsx_path)
    result = {}
    for ws in wb.worksheets:
        merges = sorted({(m.min_row, m.max_row) for m in ws.merged_cells.ranges if m.max_row > m.min_row})
        result[ws.title] = merges
    return result


def align_page_breaks(work_dir: Path, wb_name: str, sheet_names: list[str],
                       merges_by_sheet: dict[str, list[tuple[int, int]]],
                       extra_breaks: dict[str, list[int]]) -> dict[str, list[int]]:
    """把落在合并行组内部的分页点，迭代挪到该行组的第一行。返回每个 sheet 最终的分页行号。"""
    final_breaks: dict[str, list[int]] = {}
    for sheet in sheet_names:
        merges = merges_by_sheet.get(sheet, [])
        applied: set[int] = set()
        breaks: list[int] = []
        for _ in range(60):
            out = run_osa(work_dir, "enum", AS_ENUM_BREAKS, wb_name, sheet)
            breaks = [int(x) for x in out.rstrip(",").split(",") if x.strip().isdigit()]
            target = None
            for b in breaks:
                candidates = [a for (a, z) in merges if a < b <= z]
                if candidates:
                    a = min(candidates)
                    if a not in applied and a > 1:
                        target = a
                        break
            if target is None:
                break
            run_osa(work_dir, "setbrk", AS_SET_BREAK, wb_name, sheet, str(target))
            applied.add(target)
        for row in extra_breaks.get(sheet, []):
            run_osa(work_dir, "setbrk", AS_SET_BREAK, wb_name, sheet, str(row))
            applied.add(row)
        final_breaks[sheet] = breaks
    return final_breaks


def render_preview(pdf_path: Path, out_dir: Path) -> int | None:
    """把 PDF 每页渲染成 PNG，方便转换后肉眼核对有没有断错。返回页数；swift 不可用时返回 None。"""
    if shutil.which("swift") is None:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    script_path = out_dir.parent / "render.swift"
    script_path.write_text(RENDER_SWIFT, encoding="utf-8")
    proc = subprocess.run(
        ["swift", str(script_path), str(pdf_path), str(out_dir)],
        capture_output=True, text=True, timeout=180,
    )
    if proc.returncode != 0:
        return None
    lines = [l for l in proc.stdout.strip().splitlines() if l.strip().isdigit()]
    return int(lines[-1]) if lines else None


def parse_extra_breaks(items: list[str]) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for item in items:
        if ":" not in item:
            raise ValueError(f"--extra-break 格式应为 '工作表名:行号'，收到: {item}")
        sheet, row = item.rsplit(":", 1)
        result.setdefault(sheet, []).append(int(row))
    return result


def main() -> int:
    if platform.system() != "Darwin":
        print("此脚本依赖 macOS 上的 Microsoft Excel，无法在当前系统运行。", file=sys.stderr)
        return 1

    parser = argparse.ArgumentParser(description="用 Microsoft Excel 把 xlsx 转成排版正确的 PDF")
    parser.add_argument("input", type=Path, help="输入的 .xlsx/.xlsm 文件")
    parser.add_argument("-o", "--output", type=Path, default=None, help="输出 PDF 路径，默认与输入同名同目录")
    parser.add_argument("--paper", choices=sorted(PAPER_SIZES), default="a3", help="纸张大小，默认 a3（宽表格更容易一页装下）")
    parser.add_argument("--orientation", choices=["landscape", "portrait"], default="landscape")
    parser.add_argument("--extra-break", action="append", default=[], metavar="工作表名:行号",
                         help="除自动对齐外，额外在某行插入分页点（可重复传）。用于把某个过短的尾页往前收")
    parser.add_argument("--no-render", action="store_true", help="跳过导出后的 PNG 预览渲染")
    parser.add_argument("--render-dir", type=Path, default=None, help="预览 PNG 输出目录，默认放在系统临时目录")
    args = parser.parse_args()

    src = args.input.expanduser().resolve()
    if not src.exists():
        print(f"输入文件不存在: {src}", file=sys.stderr)
        return 1
    if src.suffix.lower() not in (".xlsx", ".xlsm"):
        print("目前只支持 .xlsx / .xlsm（内部是 zip 结构，才能直接改写打印设置的 XML）。"
              "如果是 .xls，先在 Excel 里另存为 .xlsx 再转。", file=sys.stderr)
        return 1

    if shutil.which("osascript") is None:
        print("找不到 osascript，无法驱动 Microsoft Excel。", file=sys.stderr)
        return 1

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("缺少 openpyxl，请先执行: pip3 install openpyxl", file=sys.stderr)
        return 1

    output = (args.output or src.with_suffix(".pdf")).expanduser().resolve()
    extra_breaks = parse_extra_breaks(args.extra_break)

    with tempfile.TemporaryDirectory(prefix="excel2pdf-") as tmp:
        work_dir = Path(tmp)
        wb_name = f"e2p-{uuid.uuid4().hex[:8]}-{src.name}"
        patched = work_dir / wb_name
        paper_size = PAPER_SIZES[args.paper]

        print(f"[1/5] 写入打印设置（{args.orientation} / {args.paper.upper()} / 缩放到一页宽）……")
        patch_print_settings(src, patched, paper_size, args.orientation)

        print("[2/5] 在 Excel 中打开临时副本……")
        run_osa(work_dir, "open", AS_OPEN, str(patched))
        for _ in range(15):
            if run_osa(work_dir, "isopen", AS_IS_OPEN, wb_name) == "YES":
                break
            time.sleep(2)
        else:
            print("Excel 没能在预期时间内打开文件，请检查 Excel 是否弹出了对话框。", file=sys.stderr)
            return 1

        merges_by_sheet = get_vertical_merges(patched)
        sheet_names = list(merges_by_sheet.keys())

        print(f"[3/5] 对齐分页点，避免切在合并单元格中间（{len(sheet_names)} 个工作表）……")
        final_breaks = align_page_breaks(work_dir, wb_name, sheet_names, merges_by_sheet, extra_breaks)

        print(f"[4/5] 导出 PDF -> {output}")
        output.parent.mkdir(parents=True, exist_ok=True)
        run_osa(work_dir, "export", AS_EXPORT, wb_name, str(output))
        run_osa(work_dir, "close", AS_CLOSE, wb_name)

        if not output.exists():
            print(f"Excel 报告导出成功，但目标文件没有生成: {output}", file=sys.stderr)
            return 1

        page_count = None
        render_dir = None
        if not args.no_render:
            render_dir = args.render_dir or (Path(tempfile.gettempdir()) / f"excel2pdf-preview-{uuid.uuid4().hex[:8]}")
            print("[5/5] 渲染分页预览图，供人工核对……")
            page_count = render_preview(output, render_dir)
        else:
            print("[5/5] 跳过预览渲染（--no-render）")

    print()
    print(f"完成: {output}")
    if page_count is not None:
        print(f"页数: {page_count}")
    for sheet, breaks in final_breaks.items():
        if breaks:
            print(f"  {sheet}: 分页行号 {breaks}")
    if render_dir is not None and render_dir.exists():
        pngs = sorted(render_dir.glob("page*.png"), key=lambda p: int(re.search(r"\d+", p.stem).group()))
        print(f"预览图目录: {render_dir}（{len(pngs)} 张，请用 Read 工具抽查几页确认没有内容被切断）")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
